# -*- coding: utf-8 -*-
"""Konversi job sheet Markdown (subset sederhana) ke DOCX berformat rapi."""
import re
import glob
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Arial"

def set_font(run, size=10.5, bold=False, italic=False, mono=False, color=None):
    fname = "Consolas" if mono else FONT
    run.font.name = fname
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), fname)
    rFonts.set(qn("w:hAnsi"), fname)

INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\n]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")
LINK = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")

def add_rich(par, text, size=10.5, base_bold=False):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            add_rich(par, part[2:-2], size=size, base_bold=True)   # isi bold bisa bersarang `kode`
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            add_rich(par, part[1:-1], size=size)
            par.runs[-1].font.italic = True
        elif part.startswith("`") and part.endswith("`"):
            set_font(par.add_run(part[1:-1]), size=size - 1, mono=True, color=(0x8B, 0x00, 0x00))
        elif part.startswith("["):
            teks, url = LINK.match(part).groups()
            set_font(par.add_run(teks), size=size, color=(0x1F, 0x4E, 0xA0))
            par.runs[-1].font.underline = True
            set_font(par.add_run(f" ({url})"), size=size - 1.5, color=(0x66, 0x66, 0x66))
        else:
            set_font(par.add_run(part), size=size, bold=base_bold)

def shade(par, hex_color="F2F2F2"):
    pPr = par._element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    pPr.append(shd)

def convert(md_path, docx_path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    doc.styles["Normal"].font.name = FONT
    doc.styles["Normal"].font.size = Pt(10.5)

    lines = open(md_path, encoding="utf-8").read().splitlines()
    in_code = False

    for line in lines:
        if line.strip().startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            p = doc.add_paragraph()          # satu paragraf per baris kode
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            shade(p, "F2F2F2")
            set_font(p.add_run(line if line else " "), size=9, mono=True)
            continue

        s = line.strip()
        if not s:
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)

        if s == "---":
            continue
        if s.startswith("# "):
            add_rich(p, s[2:], size=15, base_bold=True)
        elif s.startswith("## "):
            add_rich(p, s[3:], size=12.5, base_bold=True)
            p.paragraph_format.space_before = Pt(10)
        elif s.startswith("### "):
            add_rich(p, s[4:], size=11, base_bold=True)
            p.paragraph_format.space_before = Pt(8)
        elif s.startswith("- [ ]"):
            add_rich(p, "\u2610  " + s[5:].strip(), size=10.5)
        elif s.startswith(("- ", "* ")):
            p.paragraph_format.left_indent = Cm(0.6)
            add_rich(p, "\u2022  " + s[2:])
        elif re.match(r"^\d+\.\s", s):
            p.paragraph_format.left_indent = Cm(0.6)
            add_rich(p, s)
        elif s.startswith(">"):
            p.paragraph_format.left_indent = Cm(0.6)
            add_rich(p, s.lstrip("> ").strip(), size=10, base_bold=False)
            for r in p.runs:
                r.font.italic = True
        else:
            add_rich(p, s)

    doc.save(docx_path)
    print("ok:", docx_path)

if __name__ == "__main__":
    os.makedirs("jobsheet/docx", exist_ok=True)
    for md in sorted(glob.glob("jobsheet/*.md")):
        nama = os.path.splitext(os.path.basename(md))[0]
        convert(md, f"jobsheet/docx/Job Sheet - {nama.replace('_', ' ').title()}.docx")
